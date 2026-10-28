
from docx import Document
from docx.shared import Inches

doc = Document()
section = doc.sections[0]
section.page_width = Inches(8.5)
section.page_height = Inches(11)
section.left_margin = Inches(1)
section.right_margin = Inches(1)
section.top_margin = Inches(1)
section.bottom_margin = Inches(1)

#styles
style_title = doc.styles.add_style("BFSC_Title", WD_STYLE_TYPE.PARAGRAPH)
style_title.font.name = "Times New Roman"
style_title.font.size = Pt(12)
style_title.font.bold = True
style_title.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.CENTER


#insert title main
line1 = doc.add_paragraph("IN THE SUPERIOR COURT OF THE STATE OF WASHINGTON", level=0, style_title="BFSC_Title")


#title second line
line2 = doc.add_paragraph("IN AND FOR BENTON COUNTY", level=0, style_title="BFSC_Title")

